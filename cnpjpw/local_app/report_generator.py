"""
Módulo para Geração de Dossiês Consolidados em PDF (reportlab) e Excel (openpyxl).
Compila métricas de risco, quadro societário, empresas vinculadas, endereços compartilhados e anotações.
"""
import io
import re
import html
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def build_argumentative_dossier(
    root_data: dict,
    all_companies: list,
    socios_list: list,
    risk_info: dict,
    shared_addresses: dict,
    ubos: list,
    notes: str = ""
) -> dict:
    """
    Gera a lógica argumentativa estruturada, jurídica e pericial sobre o cluster societário.
    Articula indícios de grupo econômico de fato, confusão patrimonial, promiscuidade operacional,
    blindagem societária e requisitos para desconsideração da personalidade jurídica (Art. 50 do Código Civil).
    """
    root_name = root_data.get('nome_empresarial') or root_data.get('razao_social') or 'EMPRESA ALVO'
    root_cnpj = root_data.get('cnpj') or root_data.get('cnpj_basico') or 'CNPJ NÃO INFORMADO'
    root_sit = (root_data.get('situacao_cadastral_descricao') or root_data.get('situacao_cadastral') or 'ATIVA').upper()

    total_comps = len(all_companies)
    total_socios = len(socios_list)
    risk_level = risk_info.get('risk_level', 'BAIXO')
    risk_score = risk_info.get('risk_score', 0)

    # Beneficiários Finais (UBO)
    ubos_nomes = [u.get('nome') for u in ubos if u.get('nome')]
    ubos_str = ", ".join(ubos_nomes) if ubos_nomes else "Controle pulverizado entre os sócios diretos"

    # Sócios PJ (Holdings interpostas)
    pj_socios = []
    pf_socios = []
    for s in socios_list:
        doc = ''.join(filter(str.isdigit, str(s.get('cnpj_cpf') or s.get('cpf_cnpj') or s.get('doc') or '')))
        nome = s.get('nome') or ''
        if len(doc) == 14 or s.get('identificador_entidade_descricao') == 'PESSOA JURIDICA':
            pj_socios.append(nome)
        else:
            if nome:
                pf_socios.append(nome)

    # Empresas irregulares no cluster
    irreg_comps = []
    for c in all_companies:
        sit = str(c.get('situacao_cadastral_descricao') or c.get('situacao_cadastral') or '').upper()
        if sit in ('INAPTA', 'BAIXADA', 'SUSPENSA', 'NULA'):
            c_name = c.get('nome_empresarial') or c.get('razao_social') or c.get('cnpj') or 'Empresa'
            irreg_comps.append(f"{c_name} ({sit})")

    # Endereços compartilhados
    shared_clusters = []
    for addr_key, addr_info in shared_addresses.items():
        comps = addr_info.get('companies', [])
        lbl = addr_info.get('label') or addr_key
        if len(comps) >= 2:
            shared_clusters.append((lbl, len(comps)))

    # 1. Inferência de Coordenação e Hipótese de Grupo Econômico
    tese_grupo = (
        f"A correlação da malha societária indica a conformação relacional de grupo econômico "
        f"em torno de {root_name} (CNPJ: {root_cnpj}), integrando {total_comps} pessoas jurídicas "
        f"e {total_socios} pessoas físicas/jurídicas qualificadas no QSA. A identidade quantitativa de integrantes do corpo "
        f"diretivo subsidia a inferência técnica de coordenação administrativa e comunhão de interesses operacionais, "
        f"constituindo hipótese investigativa de atuação sob direção comum."
    )

    # 2. Correlação Espacial de Domicílios Fiscais
    if shared_clusters:
        addrs_desc = "; ".join([f"Logradouro '{lbl}' com registro simultâneo de {qtd} CNPJs" for lbl, qtd in shared_clusters[:3]])
        arg_promiscuidade = (
            f"Registrou-se a coincidência de domicílio fiscal entre entidades formalmente autônomas: "
            f"constam {len(shared_clusters)} estabelecimentos com multiplicidade cadastral de 2 ou mais pessoas jurídicas, "
            f"destacando-se: {addrs_desc}. A sobreposição cadastral de múltiplos CNPJs no mesmo espaço geográfico, "
            f"sem evidência registral de segregação de instalações físicas, subsidia a suposição técnica de confusão patrimonial "
            f"e compartilhamento operacional, a ser constatada in loco."
        )
    else:
        arg_promiscuidade = (
            f"As pessoas jurídicas mapeadas na malha societária possuem endereços cadastrais formalmente distintos. "
            f"Sugere-se averiguação da correspondência empírica das sedes frente aos registros perante os órgãos fazendários."
        )

    # 3. Estruturação em Camadas & Beneficiários Finais (UBO)
    if pj_socios:
        pj_str = ", ".join(pj_socios[:3])
        arg_blindagem = (
            f"Constatou-se a interposição de {len(pj_socios)} pessoa(s) jurídica(s) na composição do QSA ({pj_str}). "
            f"A análise da cadeia de controle societário aponta a convergência da titularidade econômica final (UBO) para: {ubos_str}. "
            f"A presença de camadas societárias intermediárias fundamenta a inferência técnica de fracionamento de titularidade, "
            f"demandando apuração da linha direta de benefício econômico."
        )
    else:
        arg_blindagem = (
            f"O quadro de sócios apresenta controle direto por pessoas físicas, convergindo a administração formal e o "
            f"benefício econômico prioritariamente para: {ubos_str}."
        )

    # 4. Correlação Temporal e Cadastral
    if irreg_comps:
        irreg_str = "; ".join(irreg_comps[:4])
        arg_irregularidade = (
            f"Mapeou-se a coexistência de {len(irreg_comps)} entidade(s) com situação cadastral inativa, suspensa ou baixada correlacionadas aos mesmos administradores: {irreg_str}. "
            f"A concomitância temporal entre cadastros baixados/inaptos e empresas plenamente ativas sob a mesma gestão "
            f"fundamenta a hipótese investigativa de descontinuidade seletiva de passivos e sucessão de fato para apuração probatória."
        )
    else:
        arg_irregularidade = (
            f"Não foram identificadas baixas compulsórias ou declarações de inaptidão no núcleo sob análise. "
            f"Mantém-se a regularidade cadastral formal aparente perante a administração tributária."
        )

    # 5. Hipótese de Subsunção Normativa (Enquadramento Técnico)
    arg_juridico = (
        f"A correlação dos indicadores quantitativos apurados (Score de Risco: {risk_score} pts | Classificação: {risk_level} | "
        f"{total_comps} sociedades correlacionadas | {total_socios} integrantes no QSA) fornece substrato empírico "
        f"para a formulação da hipótese de incidência do Art. 50 do Código Civil (com redação dada pela Lei 13.874/2019), "
        f"especificamente sob a modalidade de confusão patrimonial (§ 2º, incisos I e III), amparada na convergência de domicílios e gestão. "
        f"Subsidia-se, outrossim, o enquadramento hipotético de grupo de fato com esteio no Art. 28, § 2º do CDC, "
        f"Art. 14 da Lei 12.846/2013 e instauração do Incidente de Desconsideração da Personalidade Jurídica (CPC, arts. 133 a 137)."
    )

    # 6. Diligências Pragmáticas Sugeridas
    diligencias = [
        "1. Constatação Física In Loco: Expedição de mandado de constatação nos endereços de multiplicidade cadastral para quantificar instalações físicas, maquinário e quadro de funcionários.",
        "2. Requisição de Vínculos Financeiros (SISBAJUD): Pesquisa quantitativa de relacionamento de contas correntes e aplicações em nome das pessoas jurídicas e dos gestores mapeados.",
        "3. Rastreamento Patrimonial de Bens Móveis e Imóveis (RENAJUD / CNIB): Consulta de ativos registrados em nome de cada entidade componente do agrupamento sob investigação.",
        "4. Fluxo de Movimentação Financeira (SIMBA/COAF): Apuração técnica de eventual trânsito circular de recursos financeiros entre as pessoas jurídicas correlacionadas.",
        "5. Intimação Contábil-Fiscal: Solicitação dos livros contábeis (ECD/ECF) para verificação documental de mútuos intercompany e segregação patrimonial."
    ]

    # Texto Integral Consolidado
    paragrafos = [
        f"=== LAUDO TÉCNICO-INVESTIGATIVO & SÍNTESE RELACIONAL ===",
        f"Alvo Central: {root_name} | CNPJ: {root_cnpj} | Situação Cadastral: {root_sit}",
        f"Data da Síntese: {datetime.now().strftime('%d/%m/%Y às %H:%M')}",
        "",
        f"1. DA CORRELAÇÃO DE SOCIEDADES E HIPÓTESE DE GRUPO ECONÔMICO:",
        tese_grupo,
        "",
        f"2. DA SOBREPOSIÇÃO DE DOMICÍLIOS FISCAIS E CONFUSÃO PATRIMONIAL HIPOTÉTICA:",
        arg_promiscuidade,
        "",
        f"3. DA ESTRUTURAÇÃO EM CAMADAS E BENEFICIÁRIOS FINAIS (UBO):",
        arg_blindagem,
        "",
        f"4. DA ASSIMETRIA TEMPORAL E CORRELAÇÕES CADASTRAIS:",
        arg_irregularidade,
        "",
        f"5. DO ENQUADRAMENTO TÉCNICO-NORMATIVO HIPOTÉTICO (ART. 50 CC E CORRELATOS):",
        arg_juridico,
        "",
        f"6. PLANO DE DILIGÊNCIAS PRAGMÁTICAS RECOMENDADAS:",
        "\n".join(diligencias)
    ]

    if notes and notes.strip():
        paragrafos.extend([
            "",
            f"7. APONTAMENTOS ESPECÍFICOS DO INVESTIGADOR:",
            notes.strip()
        ])

    texto_integral = "\n".join(paragrafos)

    return {
        "tese_grupo": tese_grupo,
        "arg_promiscuidade": arg_promiscuidade,
        "arg_blindagem": arg_blindagem,
        "arg_irregularidade": arg_irregularidade,
        "arg_juridico": arg_juridico,
        "diligencias": diligencias,
        "texto_integral": texto_integral
    }


def generate_excel_dossier(
    root_data: dict,
    all_companies: list,
    socios_list: list,
    risk_info: dict,
    shared_addresses: dict,
    ubos: list,
    manual_nodes: list,
    manual_edges: list,
    notes: str = ""
) -> bytes:
    """Gera uma pasta de trabalho Excel (.xlsx) com múltiplas abas formatadas."""
    wb = openpyxl.Workbook()
    
    # Estilos padrão
    header_fill = PatternFill(start_color="0D47A1", end_color="0D47A1", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="0D47A1")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    
    # ------------------ Aba 1: Resumo & Risco ------------------
    ws_resumo = wb.active
    ws_resumo.title = "Resumo Executivo"
    
    root_name = root_data.get('nome_empresarial') or root_data.get('razao_social') or ''
    root_cnpj = root_data.get('cnpj') or root_data.get('cnpj_basico') or ''

    ws_resumo.append(["DOSSIÊ DE INTELIGÊNCIA SOCIETÁRIA & COMPLIANCE"])
    ws_resumo["A1"].font = title_font
    ws_resumo.append(["Data de Emissão:", datetime.now().strftime("%d/%m/%Y %H:%M")])
    ws_resumo.append(["Empresa Central:", f"{root_name} (CNPJ: {root_cnpj})"])
    ws_resumo.append([])
    
    # Indicadores de Risco
    ws_resumo.append(["--- INDICADORES DE RISCO & CONFORMIDADE ---"])
    ws_resumo["A5"].font = bold_font
    ws_resumo.append(["Nível de Risco do Grupo:", risk_info.get('risk_level', 'BAIXO')])
    ws_resumo.append(["Score de Risco:", f"{risk_info.get('risk_score', 0)} pontos"])
    ws_resumo.append(["Situação Cadastral Central:", risk_info.get('situacao', 'ATIVA')])
    ws_resumo.append(["Empresa Recente (< 1 ano):", "Sim" if risk_info.get('is_recente') else "Não"])
    
    flags = risk_info.get('risk_flags', [])
    if flags:
        ws_resumo.append(["Alertas Detectados:"])
        for f in flags:
            ws_resumo.append(["", f"⚠️ {f}"])
    else:
        ws_resumo.append(["Alertas Detectados:", "Nenhum alerta crítico encontrado."])
        
    ws_resumo.append([])
    ws_resumo.append(["--- ANOTAÇÕES DA INVESTIGAÇÃO ---"])
    ws_resumo.append(["Notas do Analista:", notes or "Nenhuma nota inserida."])

    # ------------------ Aba 2: Empresas do Grupo ------------------
    ws_emp = wb.create_sheet(title="Empresas Vinculadas")
    headers_emp = ["CNPJ", "Razão Social", "Situação Cadastral", "CNAE Principal", "Município/UF"]
    ws_emp.append(headers_emp)
    for col_idx in range(1, len(headers_emp) + 1):
        cell = ws_emp.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        
    for emp in all_companies:
        cnpj_val = emp.get('cnpj') or emp.get('cnpj_basico') or ''
        razao = emp.get('nome_empresarial') or emp.get('razao_social') or ''
        sit = emp.get('situacao_cadastral_descricao') or 'Ativa'
        cnae = emp.get('cnae_fiscal_principal_descricao') or ''
        mun = f"{emp.get('municipio_desc', '') or ''}/{emp.get('uf', '') or ''}"
        ws_emp.append([cnpj_val, razao, sit, cnae, mun])

    # ------------------ Aba 3: Quadro de Sócios & UBO ------------------
    ws_soc = wb.create_sheet(title="Quadro Societário & UBO")
    headers_soc = ["Nome do Sócio", "Documento", "Qualificação", "Tipo Entidade", "Beneficiário Final (UBO)"]
    ws_soc.append(headers_soc)
    for col_idx in range(1, len(headers_soc) + 1):
        cell = ws_soc.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font

    ubo_nomes = {u.get('nome') for u in ubos}
    for s in socios_list:
        nome_s = s.get('nome') or ''
        doc_s = s.get('cnpj_cpf') or ''
        qualif_s = s.get('qualificacao_descricao') or ''
        tipo_s = s.get('identificador_entidade_descricao') or ''
        is_ubo = "Sim (UBO)" if nome_s in ubo_nomes else "Sócio Direto"
        ws_soc.append([nome_s, doc_s, qualif_s, tipo_s, is_ubo])

    # ------------------ Aba 4: Endereços Compartilhados ------------------
    ws_end = wb.create_sheet(title="Endereços Compartilhados")
    headers_end = ["Endereço Compartilhado", "Qtd Empresas", "CNPJs no Endereço"]
    ws_end.append(headers_end)
    for col_idx in range(1, len(headers_end) + 1):
        cell = ws_end.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font

    for addr_key, addr_info in shared_addresses.items():
        label_addr = addr_info.get('label', '')
        comps = addr_info.get('companies', [])
        cnpjs_str = ", ".join([c.get('cnpj') or c.get('cnpj_basico') or '' for c in comps])
        ws_end.append([label_addr, len(comps), cnpjs_str])

    # ------------------ Aba 5: Lógica Argumentativa ------------------
    arg_res = build_argumentative_dossier(
        root_data=root_data,
        all_companies=all_companies,
        socios_list=socios_list,
        risk_info=risk_info,
        shared_addresses=shared_addresses,
        ubos=ubos,
        notes=notes
    )

    ws_arg = wb.create_sheet(title="Lógica Argumentativa")
    headers_arg = ["Eixo Investigativo / Dimensão", "Síntese Argumentativa & Subsunção Fática", "Enquadramento Legal / Diligências"]
    ws_arg.append(headers_arg)
    for col_idx in range(1, len(headers_arg) + 1):
        cell = ws_arg.cell(row=1, column=col_idx)
        cell.fill = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
        cell.font = header_font

    ws_arg.append(["1. Coordenação e Grupo Econômico Hipotético", arg_res.get("tese_grupo", ""), "Art. 2º, § 2º da CLT / Art. 265 da Lei 6.404/76"])
    ws_arg.append(["2. Coincidência Espacial de Domicílios Fiscais", arg_res.get("arg_promiscuidade", ""), "Art. 50, § 2º, I e III do Código Civil / Confusão Patrimonial"])
    ws_arg.append(["3. Estruturação em Camadas e UBO", arg_res.get("arg_blindagem", ""), "Art. 50 do CC / Rastreamento de Beneficiário Final (IN RFB 2.119/2022)"])
    ws_arg.append(["4. Correlação Temporal e Assimetria Cadastral", arg_res.get("arg_irregularidade", ""), "Hipótese de sucessão empresarial de fato / Art. 1.146 do CC"])
    ws_arg.append(["5. Subsunção Normativa Hipotética", arg_res.get("arg_juridico", ""), "Art. 50 do CC / Art. 28 do CDC / CPC arts. 133 a 137"])
    ws_arg.append(["6. Diligências Pragmáticas Recomendadas", "\n".join(arg_res.get("diligencias", [])), "Constatação In Loco, SISBAJUD, RENAJUD, CNIB e SIMBA"])

    if notes and notes.strip():
        ws_arg.append(["7. Anotações Adicionais do Investigador", notes.strip(), "Parecer individualizado do analista"])

    for row in ws_arg.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Auto-ajuste da largura das colunas em todas as abas
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_pdf_dossier(
    root_data: dict,
    all_companies: list[dict],
    socios_list: list[dict],
    risk_info: dict,
    shared_addresses: dict,
    ubos: list[dict],
    notes: str = ""
) -> bytes:
    """Gera um relatório formal em PDF com ReportLab."""
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0D47A1"),
        spaceAfter=6
    )
    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#546E7A"),
        spaceAfter=14
    )
    sec_style = ParagraphStyle(
        'DocSec',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#1565C0"),
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#212121")
    )
    alert_style = ParagraphStyle(
        'DocAlert',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#C62828"),
        bold=True
    )

    story = []

    # Cabeçalho
    story.append(Paragraph("DOSSIÊ DE INTELIGÊNCIA & COMPLIANCE SOCIETÁRIO", title_style))
    dt_str = datetime.now().strftime("%d/%m/%Y às %H:%M")
    story.append(Paragraph(f"Emitido em: {dt_str} | Análise Automatizada de Relacionamentos CNPJ", sub_style))

    # Resumo da Empresa Central
    root_name = root_data.get('nome_empresarial') or root_data.get('razao_social') or 'Empresa'
    root_cnpj = root_data.get('cnpj') or root_data.get('cnpj_basico') or ''
    sit_central = root_data.get('situacao_cadastral_descricao') or 'Ativa'
    
    info_table_data = [
        [Paragraph("<b>Empresa Central:</b>", body_style), Paragraph(root_name, body_style)],
        [Paragraph("<b>CNPJ:</b>", body_style), Paragraph(root_cnpj, body_style)],
        [Paragraph("<b>Situação Cadastral:</b>", body_style), Paragraph(sit_central, body_style)],
        [Paragraph("<b>CNAE Principal:</b>", body_style), Paragraph(root_data.get('cnae_fiscal_principal_descricao', 'N/A'), body_style)]
    ]
    t_info = Table(info_table_data, colWidths=[130, 410])
    t_info.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F5F7FA")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CFD8DC")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t_info)
    story.append(Spacer(1, 10))

    # Seção de Lógica Argumentativa e Parecer Técnico
    arg_res = build_argumentative_dossier(
        root_data=root_data,
        all_companies=all_companies,
        socios_list=socios_list,
        risk_info=risk_info,
        shared_addresses=shared_addresses,
        ubos=ubos,
        notes=notes
    )

    story.append(Paragraph("Síntese Técnico-Pericial & Inferências Investigativas", sec_style))
    p_arg_intro = Paragraph(
        "A presente síntese técnico-investigativa consolida correlações cadastrais, societárias e espaciais "
        "apuradas na rede de relacionamentos, estruturando hipóteses fáticas para instrução probatória "
        "e eventual subsunção às hipóteses normativas do Art. 50 do Código Civil e medidas de constrição patrimonial.",
        body_style
    )
    story.append(p_arg_intro)
    story.append(Spacer(1, 6))

    arg_rows = [
        [Paragraph("<b>1. Coordenação & Grupo Econômico:</b>", body_style), Paragraph(arg_res.get("tese_grupo", ""), body_style)],
        [Paragraph("<b>2. Coincidência Espacial de Domicílios:</b>", body_style), Paragraph(arg_res.get("arg_promiscuidade", ""), body_style)],
        [Paragraph("<b>3. Estruturação em Camadas & UBO:</b>", body_style), Paragraph(arg_res.get("arg_blindagem", ""), body_style)],
        [Paragraph("<b>4. Correlação Temporal & Cadastral:</b>", body_style), Paragraph(arg_res.get("arg_irregularidade", ""), body_style)],
        [Paragraph("<b>5. Subsunção Normativa Hipotética:</b>", body_style), Paragraph(arg_res.get("arg_juridico", ""), body_style)],
        [Paragraph("<b>6. Diligências Pragmáticas Sugeridas:</b>", body_style), Paragraph("<br/>".join(arg_res.get("diligencias", [])), body_style)],
    ]

    t_arg = Table(arg_rows, colWidths=[140, 400])
    t_arg.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F0F4F8")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#B0BEC5")),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t_arg)
    story.append(Spacer(1, 10))

    # Quadro Societário
    story.append(Paragraph("Quadro Societário & Beneficiários Finais (UBO)", sec_style))
    soc_table_data = [["Nome", "Documento", "Qualificação", "Tipo"]]
    ubo_nomes = {u.get('nome') for u in ubos}

    for s in socios_list[:15]:  # Primeiros 15 para não estourar
        nome = s.get('nome') or ''
        if nome in ubo_nomes:
            nome = f"👑 {nome} (UBO)"
        soc_table_data.append([
            Paragraph(nome[:30], body_style),
            Paragraph(s.get('cnpj_cpf', '') or '-', body_style),
            Paragraph((s.get('qualificacao_descricao') or '')[:20], body_style),
            Paragraph((s.get('identificador_entidade_descricao') or '')[:15], body_style)
        ])

    t_soc = Table(soc_table_data, colWidths=[200, 100, 140, 100])
    t_soc.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0D47A1")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E0E0E0")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(t_soc)
    story.append(Spacer(1, 10))

    # Endereços Compartilhados
    if shared_addresses:
        story.append(Paragraph("Endereços Compartilhados (Possível Cluster / Fachada)", sec_style))
        end_data = [["Endereço", "Empresas Vinculadas"]]
        for _, a_info in list(shared_addresses.items())[:5]:
            lbl = a_info.get('label', '')
            qtd = len(a_info.get('companies', []))
            end_data.append([Paragraph(lbl[:50], body_style), Paragraph(f"{qtd} empresas no mesmo local", body_style)])
        t_end = Table(end_data, colWidths=[360, 180])
        t_end.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#00838F")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E0E0E0")),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(t_end)
        story.append(Spacer(1, 10))

    # Notas do Analista & Parecer Pericial
    if notes:
        import html
        import re
        story.append(Paragraph("Parecer Técnico & Embasamento Investigativo", sec_style))
        story.append(Spacer(1, 4))
        for line in notes.split('\n'):
            line_str = line.strip()
            if not line_str:
                story.append(Spacer(1, 3))
                continue
            
            # Escapar caracteres HTML/XML para proteger o parser do ReportLab
            safe_text = html.escape(line_str)
            # Converter **negrito** para <b>negrito</b>
            safe_text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', safe_text)
            
            if safe_text.startswith('#'):
                header_text = safe_text.lstrip('#').strip()
                story.append(Spacer(1, 5))
                story.append(Paragraph(f"<b>{header_text}</b>", body_style))
            elif safe_text.startswith('- ') or safe_text.startswith('* '):
                bullet_content = safe_text[2:].strip()
                story.append(Paragraph(f"&bull; {bullet_content}", body_style))
            else:
                story.append(Paragraph(safe_text, body_style))
        story.append(Spacer(1, 10))

    doc.build(story)
    return output.getvalue()


def generate_compiled_cards_pdf(
    nodes_list: list[dict],
    edges_list: list[dict] = None,
    all_companies: list[dict] = None,
    root_data: dict = None,
    ubos: list = None
) -> bytes:
    """
    Compila os cartões de CNPJ (Pessoas Jurídicas) e Fichas Cadastrais (Pessoas Físicas)
    de todas as entidades presentes no Grafo em um único arquivo PDF.
    
    Cada entidade possui sua própria página/seção dedicada com formatação profissional,
    dados cadastrais completos, quadro societário (QSA) e vínculos identificados na rede.
    """
    import html as html_lib

    output = io.BytesIO()
    pdf_doc = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=28,
        leftMargin=28,
        topMargin=28,
        bottomMargin=28
    )

    styles = getSampleStyleSheet()

    # Estilos padronizados
    title_main_style = ParagraphStyle(
        'CardMainTitle',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0D47A1"),
        alignment=1,  # Centered
        spaceAfter=4
    )
    title_sub_style = ParagraphStyle(
        'CardSubTitle',
        parent=styles['Normal'],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#546E7A"),
        alignment=1,
        spaceAfter=12
    )
    section_hdr_style = ParagraphStyle(
        'CardSecHeader',
        parent=styles['Heading2'],
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#1565C0"),
        spaceBefore=8,
        spaceAfter=4
    )
    card_title_style = ParagraphStyle(
        'CardOfficialTitle',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#1A237E"),
        alignment=1,
        bold=True
    )
    card_sub_official = ParagraphStyle(
        'CardSubOfficial',
        parent=styles['Normal'],
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#37474F"),
        alignment=1
    )
    flabel_style = ParagraphStyle(
        'FieldLabel',
        parent=styles['Normal'],
        fontSize=6.5,
        leading=8,
        textColor=colors.HexColor("#546E7A"),
        bold=True
    )
    fval_style = ParagraphStyle(
        'FieldValue',
        parent=styles['Normal'],
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#111111"),
        bold=True
    )
    fval_norm_style = ParagraphStyle(
        'FieldValueNorm',
        parent=styles['Normal'],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#263238")
    )
    table_hdr_cell = ParagraphStyle(
        'TableHdrCell',
        parent=styles['Normal'],
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white,
        bold=True
    )
    table_row_cell = ParagraphStyle(
        'TableRowCell',
        parent=styles['Normal'],
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#212121")
    )

    story = []

    # Mapas auxiliares para enriquecimento
    ubo_names = {str(u.get('nome', '')).strip().lower() for u in (ubos or []) if isinstance(u, dict) and u.get('nome')}
    
    company_lookup = {}
    if all_companies:
        for c in all_companies:
            if not isinstance(c, dict):
                continue
            raw_c = str(c.get('cnpj') or c.get('cnpj_basico') or '')
            digs = re.sub(r'\D', '', raw_c)
            if digs:
                company_lookup[digs] = c
                if len(digs) >= 8:
                    company_lookup[digs[:8]] = c
    if root_data and isinstance(root_data, dict):
        raw_r = str(root_data.get('cnpj') or root_data.get('cnpj_basico') or '')
        digs_r = re.sub(r'\D', '', raw_r)
        if digs_r:
            company_lookup[digs_r] = root_data
            if len(digs_r) >= 8:
                company_lookup[digs_r[:8]] = root_data

    # Mapeamento de arestas para encontrar empresas vinculadas a sócios
    socio_to_companies = {}
    edges_list = edges_list or []
    node_id_to_comp = {}

    for n in nodes_list:
        nt = str(n.get('_type') or n.get('type') or '').upper()
        if 'EMPRESA' in nt:
            c_det = n.get('_details') or {}
            c_label = n.get('_raw_label') or n.get('label') or n.get('id')
            c_cnpj = c_det.get('cnpj') or n.get('_raw_val') or ''
            node_id_to_comp[n.get('id')] = {
                'id': n.get('id'),
                'razao_social': c_det.get('razao_social') or c_label,
                'cnpj': c_cnpj,
                'situacao': c_det.get('situacao_cadastral') or 'ATIVA'
            }

    for e in edges_list:
        src = e.get('from')
        dst = e.get('to')
        lbl = str(e.get('label') or 'Sócio').strip()

        # Verifica se src é empresa e dst é sócio
        if src in node_id_to_comp and dst not in node_id_to_comp:
            socio_to_companies.setdefault(dst, []).append({
                'empresa': node_id_to_comp[src]['razao_social'],
                'cnpj': node_id_to_comp[src]['cnpj'],
                'situacao': node_id_to_comp[src]['situacao'],
                'vinculo': lbl
            })
        elif dst in node_id_to_comp and src not in node_id_to_comp:
            socio_to_companies.setdefault(src, []).append({
                'empresa': node_id_to_comp[dst]['razao_social'],
                'cnpj': node_id_to_comp[dst]['cnpj'],
                'situacao': node_id_to_comp[dst]['situacao'],
                'vinculo': lbl
            })

    # Separação e Classificação das Entidades do Grafo
    pj_nodes = []
    pf_nodes = []
    seen_ids = set()

    # 1. PJs
    for n in nodes_list:
        nid = n.get('id')
        if nid in seen_ids:
            continue
        nt = str(n.get('_type') or n.get('type') or '').upper()
        if 'EMPRESA' in nt or (n.get('_is_accountant') and 'EMPRESA' in str(n.get('_details', {}).get('tipo_entidade', '')).upper()):
            seen_ids.add(nid)
            is_root = (nt == 'EMPRESA_ROOT' or n.get('_is_root') or nid == 'empresa_root')
            pj_nodes.append((0 if is_root else 1, n))

    pj_nodes.sort(key=lambda x: (x[0], str(x[1].get('_details', {}).get('razao_social') or x[1].get('label') or '')))
    pj_nodes = [item[1] for item in pj_nodes]

    # 2. PFs
    for n in nodes_list:
        nid = n.get('id')
        if nid in seen_ids:
            continue
        nt = str(n.get('_type') or n.get('type') or '').upper()
        if nt in ('SOCIO', 'UBO') or ('SOCIO' in nt) or ('UBO' in nt):
            seen_ids.add(nid)
            s_name = str(n.get('_raw_val') or n.get('label') or '').strip().lower()
            is_ubo = (nt == 'UBO' or s_name in ubo_names)
            pf_nodes.append((0 if is_ubo else 1, n))

    pf_nodes.sort(key=lambda x: (x[0], str(x[1].get('_details', {}).get('nome') or x[1].get('label') or '')))
    pf_nodes = [item[1] for item in pf_nodes]

    total_pjs = len(pj_nodes)
    total_pfs = len(pf_nodes)
    total_entidades = total_pjs + total_pfs

    # Helper de formatação de CNPJ
    def fmt_cnpj(digits_or_str: str) -> str:
        d = re.sub(r'\D', '', str(digits_or_str or ''))
        if len(d) == 14:
            return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:14]}"
        return digits_or_str or '-'

    # Helper de formatação de CPF
    def fmt_doc(doc_str: str) -> str:
        d = str(doc_str or '').strip()
        if not d:
            return "Não informado"
        digs = re.sub(r'\D', '', d)
        if len(digs) == 11 and d.isdigit():
            return f"***.{digs[3:6]}.{digs[6:9]}-**"
        return d

    # =========================================================================
    # PÁGINA 1: CAPA & ÍNDICE CONSOLIDADO DAS ENTIDADES DO GRAFO
    # =========================================================================
    story.append(Paragraph("CARTÕES CADASTRAIS CONSOLIDADOS DA REDE", title_main_style))
    story.append(Paragraph("Compilação de Comprovantes Cadastrais (Pessoas Jurídicas e Físicas) do Grafo", title_sub_style))

    root_label = ""
    root_cnpj_val = ""
    if root_data:
        root_label = root_data.get('nome_empresarial') or root_data.get('razao_social') or ''
        root_cnpj_val = fmt_cnpj(root_data.get('cnpj') or root_data.get('cnpj_basico') or '')

    dt_str = datetime.now().strftime("%d/%m/%Y às %H:%M")
    
    meta_data = [
        [Paragraph("<b>Empresa Central Investigada:</b>", flabel_style), Paragraph(f"<b>{root_label}</b> (CNPJ: {root_cnpj_val})", fval_norm_style)],
        [Paragraph("<b>Data da Compilação:</b>", flabel_style), Paragraph(dt_str, fval_norm_style)],
        [Paragraph("<b>Quantitativo de Entidades na Rede:</b>", flabel_style), Paragraph(f"<b>{total_entidades} entidades</b> ({total_pjs} Pessoas Jurídicas • {total_pfs} Pessoas Físicas / Sócios)", fval_norm_style)],
        [Paragraph("<b>Finalidade do Documento:</b>", flabel_style), Paragraph("Instrução probatória, comprovação de vínculos societários e análise cadastral individualizada.", fval_norm_style)]
    ]
    t_meta = Table(meta_data, colWidths=[160, 396])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CFD8DC")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 14))

    # Tabela Índice de Entidades
    story.append(Paragraph("Índice de Entidades Contidas neste Volume", section_hdr_style))
    
    idx_headers = [
        Paragraph("<b>#</b>", table_hdr_cell),
        Paragraph("<b>Tipo</b>", table_hdr_cell),
        Paragraph("<b>Nome Empresarial / Nome Completo</b>", table_hdr_cell),
        Paragraph("<b>Documento (CNPJ / CPF)</b>", table_hdr_cell),
        Paragraph("<b>Condição / Situação</b>", table_hdr_cell)
    ]
    idx_data = [idx_headers]

    seq = 1
    for n in pj_nodes:
        d = n.get('_details') or {}
        razao = d.get('razao_social') or n.get('_raw_label') or n.get('label') or n.get('id')
        cnpj_f = d.get('cnpj') or fmt_cnpj(n.get('_raw_val'))
        sit = d.get('situacao_cadastral') or 'ATIVA'
        is_root = 'EMPRESA_ROOT' in str(n.get('_type') or '').upper() or n.get('_is_root')
        tipo_lbl = "🏢 PJ (Raiz)" if is_root else "🏢 PJ"
        idx_data.append([
            Paragraph(str(seq), table_row_cell),
            Paragraph(tipo_lbl, table_row_cell),
            Paragraph(f"<b>{razao[:45]}</b>", table_row_cell),
            Paragraph(f"<font face='Courier'>{cnpj_f}</font>", table_row_cell),
            Paragraph(sit, table_row_cell)
        ])
        seq += 1

    for n in pf_nodes:
        d = n.get('_details') or {}
        nome = d.get('nome') or n.get('_raw_val') or n.get('label') or n.get('id')
        doc_pf_val = fmt_doc(d.get('documento') or '')
        nt = str(n.get('_type') or '').upper()
        is_ubo = nt == 'UBO' or (str(nome).strip().lower() in ubo_names)
        tipo_lbl = "👑 UBO" if is_ubo else "👤 Sócio"
        qual = d.get('qualificacao') or 'Sócio'
        idx_data.append([
            Paragraph(str(seq), table_row_cell),
            Paragraph(tipo_lbl, table_row_cell),
            Paragraph(f"<b>{nome[:45]}</b>", table_row_cell),
            Paragraph(f"<font face='Courier'>{doc_pf_val}</font>", table_row_cell),
            Paragraph(qual[:25], table_row_cell)
        ])
        seq += 1

    t_idx = Table(idx_data, colWidths=[24, 60, 246, 120, 106])
    t_idx.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0D47A1")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CFD8DC")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")])
    ]))
    story.append(t_idx)

    # =========================================================================
    # CARTÕES DE PESSOAS JURÍDICAS (CARTÃO CNPJ OFICIAL)
    # =========================================================================
    for n in pj_nodes:
        story.append(PageBreak())

        d = n.get('_details') or {}
        raw_cnpj = str(d.get('cnpj') or n.get('_raw_val') or n.get('id'))
        digs = re.sub(r'\D', '', raw_cnpj)
        
        # Enriquecimento com company_lookup se disponível
        enriched = company_lookup.get(digs) or (company_lookup.get(digs[:8]) if len(digs) >= 8 else {})
        
        razao = d.get('razao_social') or enriched.get('nome_empresarial') or enriched.get('razao_social') or n.get('label') or 'Razão Social Não Informada'
        fantasia = d.get('nome_fantasia') or enriched.get('nome_fantasia') or '*****'
        cnpj_fmt = d.get('cnpj') or fmt_cnpj(digs)
        
        matriz_filial = "MATRIZ"
        if enriched.get('identificador_matriz_filial') == 2 or (len(digs) == 14 and not digs.endswith('0001')):
            matriz_filial = "FILIAL"

        data_abertura = d.get('data_abertura') or enriched.get('data_inicio_atividade') or enriched.get('data_abertura') or 'Não informada'
        situacao = (d.get('situacao_cadastral') or enriched.get('situacao_cadastral_descricao') or 'ATIVA').upper()
        data_situacao = d.get('data_situacao') or enriched.get('data_situacao_cadastral') or ''
        
        sit_color_hex = "#2E7D32" if situacao == "ATIVA" else "#C62828"
        sit_bg_hex = "#E8F5E9" if situacao == "ATIVA" else "#FFEBEE"

        cnae_cod = d.get('cnae_codigo') or enriched.get('cnae_fiscal_principal') or ''
        cnae_desc = d.get('cnae_descricao') or enriched.get('cnae_fiscal_principal_descricao') or 'Atividade não informada'
        nat_jur = d.get('natureza_juridica') or enriched.get('natureza_juridica_descricao') or 'Não informada'
        cap_soc = d.get('capital_social') or (f"R$ {enriched.get('capital_social'):,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.') if enriched.get('capital_social') else 'Não informado')

        # Endereço
        end_str = d.get('endereco') or ''
        if not end_str and enriched:
            p = []
            if enriched.get('logradouro'): p.append(str(enriched['logradouro']))
            if enriched.get('numero'): p.append(f"nº {enriched['numero']}")
            if enriched.get('bairro'): p.append(f"Bairro {enriched['bairro']}")
            if enriched.get('municipio'): p.append(f"{enriched['municipio']}/{enriched.get('uf','')}")
            if enriched.get('cep'): p.append(f"CEP: {enriched['cep']}")
            end_str = ", ".join(p)
        if not end_str:
            end_str = "Não informado"

        # Contatos
        tels_list = d.get('telefones') or []
        tels_str = " • ".join(tels_list) if tels_list else (enriched.get('ddd_telefone_1') or 'Não informado')
        emails_list = d.get('emails') or []
        emails_str = " • ".join(emails_list) if emails_list else (enriched.get('correio_eletronico') or 'Não informado')

        # CABEÇALHO DO CARTÃO CNPJ (Estilo Receita Federal)
        header_table_data = [
            [
                Paragraph("<b>REPÚBLICA FEDERATIVA DO BRASIL</b><br/><b>CADASTRO NACIONAL DA PESSOA JURÍDICA</b>", card_title_style)
            ],
            [
                Paragraph("<b>COMPROVANTE DE INSCRIÇÃO E DE SITUAÇÃO CADASTRAL</b>", card_sub_official)
            ]
        ]
        t_card_header = Table(header_table_data, colWidths=[556])
        t_card_header.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#ECEFF1")),
            ('GRID', (0, 0), (-1, -1), 1, colors.HexColor("#455A64")),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_card_header)
        story.append(Spacer(1, 2))

        # GRADE PRINCIPAL DO CARTÃO CNPJ
        card_grid_data = [
            # Linha 1: CNPJ | MATRIZ/FILIAL | DATA ABERTURA
            [
                Paragraph(f"<font color='#546E7A'><b>NÚMERO DE INSCRIÇÃO</b></font><br/><b><font size='10'>{cnpj_fmt}</font></b><br/><b>{matriz_filial}</b>", flabel_style),
                Paragraph(f"<font color='#546E7A'><b>DATA DE ABERTURA</b></font><br/><b>{data_abertura}</b>", flabel_style)
            ],
            # Linha 2: NOME EMPRESARIAL
            [
                Paragraph(f"<font color='#546E7A'><b>NOME EMPRESARIAL</b></font><br/><font size='9.5' color='#0D47A1'><b>{razao}</b></font>", flabel_style),
                ""
            ],
            # Linha 3: NOME FANTASIA
            [
                Paragraph(f"<font color='#546E7A'><b>TÍTULO DO ESTABELECIMENTO (NOME FANTASIA)</b></font><br/><b>{fantasia}</b>", flabel_style),
                ""
            ],
            # Linha 4: CNAE PRINCIPAL
            [
                Paragraph(f"<font color='#546E7A'><b>CÓDIGO E DESCRIÇÃO DA ATIVIDADE ECONÔMICA PRINCIPAL</b></font><br/><b>{cnae_cod}</b> — {cnae_desc}", flabel_style),
                ""
            ],
            # Linha 5: NATUREZA JURÍDICA
            [
                Paragraph(f"<font color='#546E7A'><b>CÓDIGO E DESCRIÇÃO DA NATUREZA JURÍDICA</b></font><br/><b>{nat_jur}</b>", flabel_style),
                ""
            ],
            # Linha 6: ENDEREÇO
            [
                Paragraph(f"<font color='#546E7A'><b>LOGRADOURO, NÚMERO, COMPLEMENTO, BAIRRO, MUNICÍPIO E UF</b></font><br/>📍 {end_str}", flabel_style),
                ""
            ],
            # Linha 7: CONTATOS
            [
                Paragraph(f"<font color='#546E7A'><b>ENDEREÇO ELETRÔNICO (E-MAIL)</b></font><br/>✉️ {emails_str}", flabel_style),
                Paragraph(f"<font color='#546E7A'><b>TELEFONE(S)</b></font><br/>📞 {tels_str}", flabel_style)
            ],
            # Linha 8: SITUAÇÃO CADASTRAL | CAPITAL SOCIAL
            [
                Paragraph(f"<font color='#546E7A'><b>SITUAÇÃO CADASTRAL</b></font><br/><font color='{sit_color_hex}' size='10'><b>{situacao}</b></font>" + (f" <font color='#546E7A'>(desde {data_situacao})</font>" if data_situacao else ""), flabel_style),
                Paragraph(f"<font color='#546E7A'><b>CAPITAL SOCIAL</b></font><br/><b>{cap_soc}</b>", flabel_style)
            ]
        ]

        t_card_grid = Table(card_grid_data, colWidths=[386, 170])
        t_card_grid.setStyle(TableStyle([
            ('SPAN', (0, 1), (1, 1)),
            ('SPAN', (0, 2), (1, 2)),
            ('SPAN', (0, 3), (1, 3)),
            ('SPAN', (0, 4), (1, 4)),
            ('SPAN', (0, 5), (1, 5)),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#78909C")),
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#FFFFFF")),
            ('BACKGROUND', (0, 7), (0, 7), colors.HexColor(sit_bg_hex)),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_card_grid)
        story.append(Spacer(1, 8))

        # QUADRO SOCIETÁRIO (QSA)
        socios_pj = d.get('socios') or enriched.get('socios') or []
        story.append(Paragraph("Quadro de Sócios e Administradores (QSA)", section_hdr_style))

        if socios_pj:
            qsa_table_data = [
                [
                    Paragraph("<b>Nome do Sócio / Administrador</b>", table_hdr_cell),
                    Paragraph("<b>Qualificação</b>", table_hdr_cell),
                    Paragraph("<b>CPF / CNPJ</b>", table_hdr_cell),
                    Paragraph("<b>Data Entrada</b>", table_hdr_cell)
                ]
            ]
            for s in socios_pj:
                if isinstance(s, dict):
                    s_nm = s.get('nome') or 'Não informado'
                    s_ql = s.get('qualificacao') or s.get('qualificacao_descricao') or s.get('qualificacao_socio_descricao') or 'Sócio'
                    s_dc = fmt_doc(s.get('doc') or s.get('cnpj_cpf') or s.get('cpf_cnpj') or s.get('cpf') or '')
                    s_dt = s.get('data_entrada') or s.get('data_entrada_sociedade') or '-'
                    qsa_table_data.append([
                        Paragraph(f"<b>{s_nm}</b>", table_row_cell),
                        Paragraph(s_ql, table_row_cell),
                        Paragraph(f"<font face='Courier'>{s_dc}</font>", table_row_cell),
                        Paragraph(str(s_dt), table_row_cell)
                    ])
            t_qsa = Table(qsa_table_data, colWidths=[240, 146, 100, 70])
            t_qsa.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1565C0")),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CFD8DC")),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")])
            ]))
            story.append(t_qsa)
        else:
            story.append(Paragraph("<i>Nenhum sócio detalhado registrado ou quadro societário administrado por terceiros.</i>", fval_norm_style))

    # =========================================================================
    # CARTÕES DE PESSOAS FÍSICAS (FICHA CADASTRAL PF & VÍNCULOS)
    # =========================================================================
    for n in pf_nodes:
        story.append(PageBreak())

        d = n.get('_details') or {}
        nome_pf = d.get('nome') or n.get('_raw_val') or n.get('label') or 'Nome Não Informado'
        doc_pf = fmt_doc(d.get('documento') or '')
        qual_pf = d.get('qualificacao') or 'Sócio'
        
        nt = str(n.get('_type') or '').upper()
        is_ubo = (nt == 'UBO' or str(nome_pf).strip().lower() in ubo_names)
        
        badge_header = "FICHA CADASTRAL DE PESSOA FÍSICA"
        if is_ubo:
            badge_header += " — 👑 BENEFICIÁRIO FINAL (UBO)"

        header_pf_data = [
            [
                Paragraph("<b>SISTEMA POMELO / CAEXLGS — REDE DE RELACIONAMENTOS</b>", card_title_style)
            ],
            [
                Paragraph(f"<b>{badge_header}</b>", card_sub_official)
            ]
        ]
        t_pf_header = Table(header_pf_data, colWidths=[556])
        t_pf_header.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#FFF8E1") if is_ubo else colors.HexColor("#ECEFF1")),
            ('GRID', (0, 0), (-1, -1), 1, colors.HexColor("#FFB300") if is_ubo else colors.HexColor("#455A64")),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_pf_header)
        story.append(Spacer(1, 4))

        # GRADE PRINCIPAL DA PESSOA FÍSICA
        grid_pf_data = [
            [
                Paragraph(f"<font color='#546E7A'><b>NOME COMPLETO</b></font><br/><font size='10.5' color='#0D47A1'><b>{nome_pf}</b></font>", flabel_style),
                Paragraph(f"<font color='#546E7A'><b>DOCUMENTO (CPF / DOC)</b></font><br/><b><font size='9.5'>{doc_pf}</font></b>", flabel_style)
            ],
            [
                Paragraph(f"<font color='#546E7A'><b>CONDICÃO NA INVESTIGAÇÃO</b></font><br/><b>{'👑 BENEFICIÁRIO FINAL (UBO)' if is_ubo else '👤 SÓCIO / ADMINISTRADOR'}</b>", flabel_style),
                Paragraph(f"<font color='#546E7A'><b>QUALIFICAÇÃO PRINCIPAL DECLARADA</b></font><br/><b>{qual_pf}</b>", flabel_style)
            ]
        ]
        t_pf_grid = Table(grid_pf_data, colWidths=[366, 190])
        t_pf_grid.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#78909C")),
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#FFFFFF")),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        story.append(t_pf_grid)
        story.append(Spacer(1, 10))

        # EMPRESAS VINCULADAS NA REDE DO GRAFO
        story.append(Paragraph("Participações Societárias & Empresas Vinculadas no Grafo", section_hdr_style))
        
        # Recupera as empresas ligadas por aresta ou por detalhes
        linked_comps = socio_to_companies.get(n.get('id'), [])
        if not linked_comps and d.get('empresas_vinculadas'):
            for ev in d.get('empresas_vinculadas'):
                linked_comps.append({
                    'empresa': ev,
                    'cnpj': 'Vinculada no Grafo',
                    'situacao': 'ATIVA',
                    'vinculo': qual_pf
                })

        if linked_comps:
            vinc_table_data = [
                [
                    Paragraph("<b>Empresa Vinculada (Razão Social)</b>", table_hdr_cell),
                    Paragraph("<b>CNPJ</b>", table_hdr_cell),
                    Paragraph("<b>Qualificação / Papel</b>", table_hdr_cell),
                    Paragraph("<b>Situação</b>", table_hdr_cell)
                ]
            ]
            for lc in linked_comps:
                vinc_table_data.append([
                    Paragraph(f"<b>{lc.get('empresa','-')}</b>", table_row_cell),
                    Paragraph(f"<font face='Courier'>{lc.get('cnpj','-')}</font>", table_row_cell),
                    Paragraph(lc.get('vinculo','Sócio'), table_row_cell),
                    Paragraph(lc.get('situacao','ATIVA'), table_row_cell)
                ])
            t_vinc = Table(vinc_table_data, colWidths=[240, 130, 116, 70])
            t_vinc.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#2E7D32") if is_ubo else colors.HexColor("#1565C0")),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CFD8DC")),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")])
            ]))
            story.append(t_vinc)
        else:
            story.append(Paragraph("<i>Nenhuma empresa adicional diretamente vinculada a este sócio nesta visualização.</i>", fval_norm_style))

    pdf_doc.build(story)
    return output.getvalue()
